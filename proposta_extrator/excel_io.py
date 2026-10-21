import os
import shutil

import openpyxl

COL_PEDIDO = 11  # K
COLS_SAIDA = {"modelo": 12, "serial": 13, "part_number": 14,
              "preco_unitario": 15, "preco_total": 16, "prazo": 17}  # L..Q


def caminho_saida(origem):
    nome, ext = os.path.splitext(origem)
    return f"{nome}_preenchido{ext}"


def abrir(origem):
    return openpyxl.load_workbook(origem, keep_vba=origem.lower().endswith(".xlsm"))


def listar_pedidos(ws, linha_inicial):
    for r in range(linha_inicial, ws.max_row + 1):
        v = ws.cell(r, COL_PEDIDO).value
        if v is not None and str(v).strip():
            yield r, v


def gravar_linha(ws, linha, dados):
    for campo, col in COLS_SAIDA.items():
        if dados.get(campo):
            ws.cell(linha, col).value = dados[campo]
